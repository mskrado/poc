import sys, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
Y=PatternFill("solid",fgColor="FFF2A8"); B=Font(bold=True)
ex = len(sys.argv)>1
wb=openpyxl.Workbook()
r=wb.active; r.title="Read me"
for row in [["CFO AI Support Payoff Worksheet"],["By Av Ledger, Enterprise ROI Desk, roiscale.ai"],[],
 ["Purpose","Test an AI customer-service business case the way Klarna's survived its SEC review: stated denominator, forecast reconciled to a booked figure, quality cost priced in."],
 ["How to use","1. Fill the yellow cells on 'Inputs' from your own logs and ledger (same 12-month period, monthly averages)."],
 ["","2. Read results on 'Model'. Payback shows 'No payback' when monthly net saving <= 0, and blank until inputs are filled."],
 ["","3. 'Klarna reference' lists only what Klarna disclosed. Build and run cost were not disclosed, so payback cannot be computed from public data."],
 ["Companion to","roiscale.ai article \"Klarna's AI Payoff Survived the SEC\""],[],["Planning aid, not financial advice."]]:
    r.append(row)
r["A1"].font=Font(bold=True,size=14); r.column_dimensions["A"].width=16; r.column_dimensions["B"].width=110
i=wb.create_sheet("Inputs"); i.append(["Input","Value","Unit","Note"])
inp=[("Period length",12,"months","Same period before and after"),
("Baseline contacts per month",20000,"contacts/month","From your logs, all channels"),
("Fully loaded human cost per contact",6,"currency/contact",""),
("AI resolution share",0.6,"% of contacts","Share handled end-to-end by AI"),
("AI run cost per AI-handled contact",0.8,"currency/contact","Run cost (licence, inference, ops)"),
("Escalation rate (AI to human)",0.1,"% of AI contacts",""),
("Cost per escalation",7,"currency/escalation",""),
("Repeat-contact rate within 7 days, before",0.1,"% of contacts",""),
("Repeat-contact rate within 7 days, after",0.12,"% of contacts",""),
("Build / one-off cost",150000,"currency","One-off"),
("Forecast saving for the period",400000,"currency","As stated in the business case"),
("Booked saving for the period",350000,"currency","From the ledger")]
for k,(n,v,u,note) in enumerate(inp,start=2):
    i.append([n,v if (ex or k==2) else None,u,note]); i.cell(k,2).fill=Y
    if "%" in u: i.cell(k,2).number_format="0.0%"
for c,w in zip("ABCD",[42,14,22,40]): i.column_dimensions[c].width=w
for c in i[1]: c.font=B
m=wb.create_sheet("Model"); m.append(["Line","Value","Formula"])
I=lambda n:f"Inputs!B{n}"
ready=f"COUNT(Inputs!B3:B11)=9"
rows=[
("Human cost per month, before",f"={I(3)}*{I(4)}","contacts x human cost"),
("Resolved contacts per month, before",f"={I(3)}*(1-{I(9)})","contacts - repeat contacts"),
("Cost per resolved contact, before",f"=IF(B3>0,B2/B3,\"\")","human cost / resolved"),
("Human cost per month, after",f"={I(3)}*(1-{I(5)})*{I(4)}","non-AI contacts x human cost"),
("AI run cost per month",f"={I(3)}*{I(5)}*{I(6)}","AI contacts x run cost"),
("Escalation cost per month",f"={I(3)}*{I(5)}*{I(7)}*{I(8)}","AI contacts x escalation rate x cost"),
("Resolved contacts per month, after",f"={I(3)}*(1-{I(10)})","contacts - repeat contacts"),
("Cost per resolved contact, after",f"=IF(B8>0,(B5+B6+B7)/B8,\"\")","(human + AI run + escalation) / resolved"),
("Gross saving per month",f"=B2-B5","human cost before - after"),
("Quality cost per month",f"={I(3)}*({I(10)}-{I(9)})*{I(4)}","change in repeat contacts x loaded cost"),
("Monthly net saving",f"=B10-B6-B7-B11","gross - run - escalation - quality"),
("Payback months",f"=IF(NOT({ready}),\"\",IF(B12<=0,\"No payback\",{I(11)}/B12))","build cost / monthly net saving"),
("Modelled net saving for the period",f"=B12*{I(2)}","monthly net x months"),
("Forecast saving",f"=IF({I(12)}=\"\",\"\",{I(12)})","from Inputs"),
("Booked saving",f"=IF({I(13)}=\"\",\"\",{I(13)})","from Inputs"),
("Gap, booked vs forecast",f"=IF(OR(B15=\"\",B16=\"\"),\"\",B16-B15)",""),
("Gap %",f"=IF(OR(B15=\"\",B16=\"\",B15=0),\"\",B17/B15)","(booked - forecast) / forecast")]
for n,f,d in rows: m.append([n,f,d])
for k in range(2,19): m.cell(k,2).number_format="#,##0.00"
m["B18"].number_format="0.0%"; m["B13"].number_format="0.0"
for c,w in zip("ABC",[38,16,44]): m.column_dimensions[c].width=w
for c in m[1]: c.font=B
k=wb.create_sheet("Klarna reference"); k.append(["Figure","Value","Source"])
P="Klarna press release, 27 Feb 2024"; S="Klarna SEC response letter, 16 Jul 2025"
for row in [("AI assistant conversations (first month)","2.3 million",P),("Share of customer service chats","Two-thirds",P),
("Work equivalent","700 full-time agents",P),("Estimated profit improvement in 2024","$40 million (forecast)",P),
("Resolution time","Under 2 min vs 11 min",P),("Repeat inquiries","-25%",P),
("Share of chats, 12 months to 31 Mar 2025","66%",S),("FTE equivalent","Over 700 (estimate)",S),
("Cost savings in 2024","About $39 million",S),("Customer satisfaction","On par with human agents",S),
("Repeat inquiries, Dec 2023 to Jan 2024","-25%",S),("Resolution time, 2024","2 min vs 12 min human average",S),
("Build / one-off cost","Not disclosed",""),("Run cost","Not disclosed",""),("Escalation cost","Not disclosed","")]: k.append(row)
k.append([]); k.append(["No build or run cost was disclosed, so payback cannot be computed from public data. No numbers here are estimates by roiscale.ai."])
for c,w in zip("ABC",[42,32,38]): k.column_dimensions[c].width=w
for c in k[1]: c.font=B
wb.move_sheet("Read me",-0)
wb.save("test.xlsx" if ex else "cfo-ai-support-payoff.xlsx")
